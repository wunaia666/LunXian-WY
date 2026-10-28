# -*- coding: utf-8 -*-
"""将十张原生动态图表加入现有客户管理工作簿，并自动核对。"""
from pathlib import Path
from datetime import datetime,date
import argparse,calendar,json,shutil,warnings
import openpyxl
import win32com.client
from 图表页面 import create_component,P
from 数据公式 import build_formulas
from 图表美化 import polish,finalize_layout
from 原生图表 import create_native_charts

warnings.filterwarnings('ignore',category=UserWarning)

def rgb(h):return int(h[:2],16)+int(h[2:4],16)*256+int(h[4:6],16)*65536

def raw_inputs(path):
    book=openpyxl.load_workbook(path,data_only=True)
    return {name:[list(row) for row in book[name].iter_rows(min_row=3,values_only=True) if row[0]] for name in ['客户档案','服务记录','回访任务']}

def month_range(year,month):
    start=date(year,month,1)
    return start,date(year+int(month==12),1 if month==12 else month+1,1)

def expected(data,year,month,store):
    start,end=month_range(year,month)
    customers={r[0]:r for r in data['客户档案']}
    in_store=lambda actual:store=='全部门店' or actual==store
    customers_count=sum(in_store(c[5]) for c in customers.values())
    services=[r for r in data['服务记录'] if isinstance(r[1],datetime) and start<=r[1].date()<end and in_store(r[4])]
    plans=completed=overdue=actual_count=0
    for r in data['回访任务']:
        if not in_store(r[3]):continue
        actual=r[11]
        done=isinstance(actual,datetime) and actual<=datetime.now()
        c=customers[r[1]]
        stopped=not actual and c[12]=='老客' and r[5] in ['新客节点','新客循环']
        if isinstance(r[9],datetime) and start<=r[9].date()<end and not stopped:
            plans+=1
            if done:completed+=1
        if not done and not stopped and not actual and isinstance(r[9],datetime) and r[9].date()<date.today():overdue+=1
        if done and start<=actual.date()<end:actual_count+=1
    weekdays=[sum(r[1].weekday()==i for r in services) for i in range(7)]
    return {'消费':sum(float(r[9] or 0) for r in services),'服务次数':len(services),'客户数':customers_count,'计划':plans,'完成':completed,'逾期':overdue,'实际完成':actual_count,'星期分布':weekdays}

def build(source,output,qa):
    source=Path(source).resolve();output=Path(output).resolve();qa=Path(qa).resolve()
    if source==output:raise ValueError('输出文件请使用新文件名，不能覆盖原始工作簿。')
    output.parent.mkdir(parents=True,exist_ok=True);qa.mkdir(parents=True,exist_ok=True)
    component=qa/'图表组件.xlsx'
    layout=create_component(component)
    formulas=build_formulas()
    temp=openpyxl.load_workbook(component,data_only=False)
    helper=temp['图表数据']
    for cell,expression in formulas.items():helper[cell]=expression
    matrix=tuple(tuple(helper.cell(r,c).value for c in range(1,11)) for r in range(1,169))
    shutil.copy2(source,output)
    inputs=raw_inputs(source)
    app=win32com.client.DispatchEx('Excel.Application')
    app.Visible=False;app.DisplayAlerts=False;app.ScreenUpdating=False
    book=parts=None
    checks=[]
    published=False
    chart_objects={}
    try:
        book=app.Workbooks.Open(str(output),UpdateLinks=0)
        app.Calculation=-4135
        if any(str(s.Name) in ['经营驾驶舱','图表分析中心','图表数据'] for s in book.Worksheets):raise ValueError('原文件已经包含同名扩展页，请选择未添加图表的原始版本。')
        parts=app.Workbooks.Open(str(component),UpdateLinks=0)
        parts.Worksheets.Copy(Before=book.Worksheets(1))
        parts.Close(False);parts=None
        data=book.Worksheets('图表数据')
        data.Range('A1:J168').Formula=matrix
        data.Range('A26:A56').NumberFormat='yyyy-mm-dd'
        data.Range('B12').NumberFormat='0.0%'
        data.Range('B16').NumberFormat='+0.0%;-0.0%;0.0%'
        # 日趋势会随月份天数变化；员工排行只保留当前范围中的人员。
        names={
            '日趋势日期':"=OFFSET('图表数据'!$G$26,0,0,DAY(EOMONTH('图表数据'!$B$2,0)),1)",
            '日趋势消费':"=OFFSET('图表数据'!$B$26,0,0,DAY(EOMONTH('图表数据'!$B$2,0)),1)",
            '日趋势回访':"=OFFSET('图表数据'!$C$26,0,0,DAY(EOMONTH('图表数据'!$B$2,0)),1)",
            '员工排行姓名':"=OFFSET('图表数据'!$G$79,0,0,'图表数据'!$J$78,1)",
            '员工排行次数':"=OFFSET('图表数据'!$H$79,0,0,'图表数据'!$J$78,1)"}
        for name,ref in names.items():book.Names.Add(Name=name,RefersTo=ref)
        chart_objects=create_native_charts(book,layout)
        for sheet_name in ['经营驾驶舱','图表分析中心']:
            ws=book.Worksheets(sheet_name)
            specs=[x for x in layout if x['sheet']==sheet_name]
            objects=ws.ChartObjects()
            assert objects.Count==len(specs),(sheet_name,objects.Count)
            for i,spec in enumerate(specs,1):
                obj=objects.Item(i);chart=obj.Chart
                region=ws.Range(spec['region']);anchor=ws.Range(spec['top_cell'])
                obj.Left=region.Left+8;obj.Top=anchor.Top+2
                obj.Width=region.Width-16;obj.Height=region.Top+region.Height-obj.Top-6
                obj.Name='图表_'+spec['key']
                obj.ShapeRange.AlternativeText=spec['title']
                chart.PlotVisibleOnly=False
                chart.HasTitle=False
                chart.ChartArea.Format.Fill.ForeColor.RGB=rgb(P['white'])
                chart.ChartArea.Format.Line.Visible=False
                chart.ChartArea.Font.Name='微软雅黑';chart.ChartArea.Font.Size=9
                chart_objects[spec['key']]=chart
                if chart.HasLegend:
                    chart.Legend.Font.Name='微软雅黑';chart.Legend.Font.Size=8
                try:
                    for axis in chart.Axes():
                        axis.TickLabels.Font.Name='微软雅黑';axis.TickLabels.Font.Size=9
                except Exception:pass
                for series in chart.SeriesCollection():
                    try:
                        if series.HasDataLabels:
                            series.DataLabels().Font.Name='微软雅黑'
                            series.DataLabels().Font.Color=rgb(P['ink'])
                    except Exception:pass
            ws.Tab.Color=rgb(P['ink'])
            ws.Activate();app.ActiveWindow.DisplayGridlines=False
            app.ActiveWindow.DisplayHeadings=False;app.ActiveWindow.Zoom=85
        qbook="'"+str(book.Name).replace("'","''")+"'"
        daily=chart_objects['daily']
        daily.SeriesCollection(1).XValues='='+qbook+'!日趋势日期'
        daily.SeriesCollection(1).Values='='+qbook+'!日趋势消费'
        daily.SeriesCollection(2).XValues='='+qbook+'!日趋势日期'
        daily.SeriesCollection(2).Values='='+qbook+'!日趋势回访'
        daily.SeriesCollection(1).Format.Fill.ForeColor.RGB=rgb(P['ink'])
        daily.SeriesCollection(2).Format.Line.ForeColor.RGB=rgb(P['orange'])
        staff=chart_objects['staff']
        staff.SeriesCollection(1).XValues='='+qbook+'!员工排行姓名'
        staff.SeriesCollection(1).Values='='+qbook+'!员工排行次数'
        data.Visible=0
        driver=book.Worksheets('经营驾驶舱')
        app.Calculation=-4105
        app.CalculateFullRebuild()
        polish(book,chart_objects)
        print('十张原生图表已添加，正在核对联动。',flush=True)
        for year,month,store in [(2026,10,'全部门店'),(2026,10,'二店'),(2027,2,'全部门店')]:
            driver.Range('B7').Value=year;driver.Range('F7').Value=month;driver.Range('J7').Value=store
            app.CalculateFullRebuild()
            for chart in chart_objects.values():chart.Refresh()
            e=expected(inputs,year,month,store)
            for field,cell in [('消费','B7'),('服务次数','B8'),('客户数','B9'),('计划','B10'),('完成','B11'),('逾期','B13'),('实际完成','B18')]:
                assert abs(float(data.Range(cell).Value)-float(e[field]))<.00001,(year,month,store,field,data.Range(cell).Value,e[field])
            assert float(driver.Range('B11').Value)==e['消费']
            assert float(driver.Range('I11').Value)==e['服务次数']
            assert float(driver.Range('P11').Value)==e['客户数']
            assert abs(sum(float(data.Range(f'B{r}').Value or 0) for r in range(26,57))-e['消费'])<.00001
            assert abs(sum(float(data.Range(f'B{r}').Value or 0) for r in range(68,75))-e['消费'])<.00001
            assert int(sum(float(data.Range(f'B{r}').Value or 0) for r in range(128,135)))==e['服务次数']
            assert [int(data.Range(f'B{r}').Value or 0) for r in range(128,135)]==e['星期分布']
            assert int(sum(float(data.Range(f'B{r}').Value or 0) for r in range(140,144)))==e['逾期']
            assert daily.SeriesCollection(1).Points().Count==calendar.monthrange(year,month)[1]
            assert staff.SeriesCollection(1).Points().Count==int(data.Range('J78').Value)
            assert float(data.Range('B12').Value)==(e['完成']/e['计划'] if e['计划'] else 0)
            checks.append({'年份':year,'月份':month,'门店':store,'核对结果':e,'趋势天数':daily.SeriesCollection(1).Points().Count,'员工图表人数':staff.SeriesCollection(1).Points().Count})
        driver.Range('B7').Value=date.today().year;driver.Range('F7').Value=date.today().month;driver.Range('J7').Value='全部门店'
        app.CalculateFullRebuild()
        # 临时录入服务记录，确认图表会随输入更新，随后原样恢复。
        services=book.Worksheets('服务记录')
        blank=next(r for r in range(3,203) if not services.Cells(r,1).Value)
        saved=services.Range(f'A{blank}:K{blank}').Formula
        old_amount=float(data.Range('B7').Value);old_count=float(data.Range('B8').Value)
        serial=(date.today()-date(1899,12,30)).days
        client_id=inputs['客户档案'][0][0]
        try:
            services.Cells(blank,1).Value='图表联动核对记录'
            services.Cells(blank,2).Value=serial
            services.Cells(blank,3).Value=client_id
            services.Cells(blank,6).Value='面部深层护理'
            services.Cells(blank,10).Value=99
            app.CalculateFullRebuild();daily.Refresh()
            assert abs(float(data.Range('B7').Value)-(old_amount+99))<.00001
            assert float(data.Range('B8').Value)==old_count+1
            assert abs(float(daily.SeriesCollection(1).Values[date.today().day-1])-float(data.Range(f'B{25+date.today().day}').Value))<.00001
            checks.append({'新增录入联动':'通过','金额变化':99,'次数变化':1,'核对记录':'核对后已恢复原始行'})
        finally:
            services.Range(f'A{blank}:K{blank}').Formula=saved
        app.CalculateFullRebuild()
        for chart in chart_objects.values():chart.Refresh()
        finalize_layout(book,chart_objects)
        driver.Activate();driver.Range('J7').Select()
        app.ActiveWindow.ScrollRow=1;app.ActiveWindow.ScrollColumn=1
        book.CheckCompatibility=False
        book.Save()
        for name in ['经营驾驶舱','图表分析中心']:
            book.Worksheets(name).ExportAsFixedFormat(0,str(qa/(name+'.pdf')))
        driver.Activate();driver.Range('J7').Select();book.Save()
        published=True
    finally:
        if parts is not None:parts.Close(False)
        if book is not None:book.Close(False)
        if published:
            # 正常重新打开并保存，清理复制组件时留下的无用链接缓存。
            normalized=app.Workbooks.Open(str(output),UpdateLinks=0)
            app.CalculateFullRebuild()
            normalized.Save();normalized.Close(False)
        app.Quit()
    original=openpyxl.load_workbook(source,data_only=False)
    final=openpyxl.load_workbook(output,data_only=False)
    cached=openpyxl.load_workbook(output,data_only=True)
    changes=[(s.title,c.coordinate) for s in original for row in s for c in row if c.value!=final[s.title][c.coordinate].value]
    errors=[(s.title,c.coordinate,c.value) for s in cached for row in s for c in row if c.data_type=='e']
    assert not changes,changes
    assert not errors,errors
    assert not final._external_links,'不应产生外部工作簿链接'
    count=sum(len(s._charts) for s in final)
    assert count==10,count
    report={'交付文件':str(output),'图表总数':count,'图表分布':{s.title:len(s._charts) for s in final if s._charts},'原表内容变动':changes,'公式错误':errors,'外部工作簿链接':len(final._external_links),'联动核对':checks,'大小':output.stat().st_size}
    (qa/'验收结果.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))
    return report

if __name__=='__main__':
    parser=argparse.ArgumentParser(description='为现有客户管理表增加经营驾驶舱和图表分析中心。')
    parser.add_argument('--source',required=True,help='现有日历工作台工作簿的绝对路径')
    parser.add_argument('--output',required=True,help='新增图表后的工作簿绝对路径')
    parser.add_argument('--qa-dir',required=True,help='预览与验收结果目录')
    args=parser.parse_args()
    build(args.source,args.output,args.qa_dir)
