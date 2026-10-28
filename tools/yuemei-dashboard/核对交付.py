# -*- coding: utf-8 -*-
"""检查旧数据保留、原生图表与跨文件改名后的动态联动。"""
from pathlib import Path
import argparse,json,shutil,warnings
import openpyxl
import win32com.client
warnings.filterwarnings('ignore',category=UserWarning)

def verify(source,target,qa_dir):
    source=Path(source).resolve();target=Path(target).resolve();qa=Path(qa_dir).resolve()
    qa.mkdir(parents=True,exist_ok=True)
    before=openpyxl.load_workbook(source,data_only=False)
    after=openpyxl.load_workbook(target,data_only=False)
    values=openpyxl.load_workbook(target,data_only=True)
    changes=[(s.title,c.coordinate) for s in before for row in s for c in row if c.value!=after[s.title][c.coordinate].value]
    errors=[(s.title,c.coordinate,c.value) for s in values for row in s for c in row if c.data_type=='e']
    assert not changes,changes
    assert not errors,errors
    assert not after._external_links
    assert after['图表数据'].sheet_state=='hidden'
    assert len(after['经营驾驶舱']._charts)==4
    assert len(after['图表分析中心']._charts)==6
    for name in ['日趋势日期','日趋势消费','日趋势回访','员工排行姓名','员工排行次数']:assert name in after.defined_names,name
    for s in before:
        assert set(s.tables)==set(after[s.title].tables),(s.title,'原始表格被修改')
        for key in s.tables:assert s.tables[key].ref==after[s.title].tables[key].ref
    renamed=qa/'转发改名核对.xlsx'
    shutil.copy2(target,renamed)
    app=win32com.client.DispatchEx('Excel.Application');app.Visible=False;app.DisplayAlerts=False
    book=None
    try:
        # 使用正常打开流程，不能依赖修复模式或原文件路径。
        book=app.Workbooks.Open(str(renamed),UpdateLinks=0,ReadOnly=True)
        assert book.LinkSources(1) is None,'转发文件不应依赖其他工作簿'
        home=book.Worksheets('经营驾驶舱')
        home.Range('B7').Value=2027;home.Range('F7').Value=2;home.Range('J7').Value='二店'
        app.CalculateFullRebuild()
        daily=home.ChartObjects().Item(1).Chart
        daily.Refresh()
        assert daily.SeriesCollection(1).Points().Count==28
        assert float(book.Worksheets('图表数据').Range('B7').Value)==0
        assert book.Worksheets('图表分析中心').ChartObjects().Count==6
        result={'图表':10,'原表变动':0,'公式错误':0,'外部链接':0,'改名转发后动态更新':'通过','二月趋势天数':28,'原生图表可编辑':'通过'}
    finally:
        if book is not None:book.Close(False)
        app.Quit()
    (qa/'交付核对.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2))
    return result

if __name__=='__main__':
    p=argparse.ArgumentParser(description='核对动态图表交付文件。')
    p.add_argument('--source',required=True);p.add_argument('--target',required=True);p.add_argument('--qa-dir',required=True)
    a=p.parse_args();verify(a.source,a.target,a.qa_dir)
