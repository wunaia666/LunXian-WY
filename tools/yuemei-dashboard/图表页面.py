# -*- coding: utf-8 -*-
"""生成可跨工作簿复制的原生图表页面与占位数据。"""
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Border,Side,Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart,LineChart,DoughnutChart,Reference
from openpyxl.chart.series import DataPoint
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.axis import ChartLines
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import Paragraph,ParagraphProperties,CharacterProperties,Font as DrawingFont
from openpyxl.chart.text import RichText
from openpyxl.utils import get_column_letter

P={'ink':'2D4A2B','sage':'7D8471','olive':'A4AC86','paper':'FAF9F6','white':'FFFFFF','line':'E1E6DD','muted':'64705F','orange':'D97757','red':'A64338','soft':'EDF1E6','gold':'936B35'}
CHARTS=[]

def area(ws,r1,c1,r2,c2,text=None,fill='white',size=11,bold=False,color='ink',align='left',number_format=None):
    for row in ws.iter_rows(min_row=r1,max_row=r2,min_col=c1,max_col=c2):
        for cell in row:
            cell.fill=PatternFill('solid',fgColor=P[fill])
            cell.font=Font(name='微软雅黑',size=size,bold=bold,color=P[color])
            cell.alignment=Alignment(horizontal=align,vertical='center',indent=1 if align=='left' else 0)
    if r1!=r2 or c1!=c2:ws.merge_cells(start_row=r1,start_column=c1,end_row=r2,end_column=c2)
    cell=ws.cell(r1,c1)
    if text is not None:cell.value=text
    if number_format:cell.number_format=number_format
    return cell

def panel(ws,r1,c1,r2,c2,title,subtitle):
    ws.row_dimensions[r1].height=17
    ws.row_dimensions[r1+1].height=13
    edge=Side(style='thin',color=P['line'])
    for row in ws.iter_rows(min_row=r1,max_row=r2,min_col=c1,max_col=c2):
        for cell in row:cell.fill=PatternFill('solid',fgColor=P['white'])
    for c in range(c1,c2+1):
        ws.cell(r1,c).border=Border(top=edge)
        ws.cell(r2,c).border=Border(bottom=edge)
    for r in range(r1,r2+1):
        ws.cell(r,c1).border=Border(left=edge)
        ws.cell(r,c2).border=Border(right=edge)
    if title:area(ws,r1,c1,r1,c2,title,size=12,bold=True)
    if subtitle:area(ws,r1+1,c1,r1+1,c2,subtitle,size=8.5,color='muted')

def rich(size=9):
    prop=CharacterProperties(latin=DrawingFont(typeface='微软雅黑'),ea=DrawingFont(typeface='微软雅黑'),sz=int(size*100),solidFill=P['ink'])
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=prop),endParaRPr=prop)])

def chart_base(chart,legend=True):
    chart.title=None
    chart.style=None
    chart.graphical_properties=GraphicalProperties(solidFill=P['white'],ln=LineProperties(noFill=True))
    chart.plot_area.spPr=GraphicalProperties(solidFill=P['white'],ln=LineProperties(noFill=True))
    chart.visible_cells_only=False
    chart.display_blanks='gap'
    chart.roundedCorners=False
    chart.txPr=rich()
    if legend:
        if chart.legend:
            chart.legend.position='b'
            chart.legend.txPr=rich(8.5)
    else:chart.legend=None
    for axis_name in ['x_axis','y_axis']:
        axis=getattr(chart,axis_name,None)
        if axis is None:continue
        axis.txPr=rich(8.5)
        axis.spPr=GraphicalProperties(ln=LineProperties(solidFill=P['line'],w=8000))
    if hasattr(chart,'y_axis'):
        chart.y_axis.axPos='l'
        chart.y_axis.delete=False
        chart.y_axis.majorGridlines=ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill=P['line'],w=7000)))
        chart.y_axis.numFmt='#,##0'
        chart.y_axis.scaling.min=0
    if hasattr(chart,'x_axis'):
        chart.x_axis.axPos='b'
        chart.x_axis.delete=False
        chart.x_axis.majorGridlines=None
    return chart

def bars(data,min_row,max_row,label_col,value_cols,horizontal=False,multi=False):
    chart=chart_base(BarChart(),legend=multi)
    chart.type='bar' if horizontal else 'col'
    chart.grouping='clustered'
    chart.overlap=0
    chart.gapWidth=80 if horizontal else 95
    for index,col in enumerate(value_cols):
        chart.add_data(Reference(data,min_col=col,min_row=min_row-1,max_row=max_row),titles_from_data=True)
        series=chart.series[-1]
        series.graphicalProperties.solidFill=[P['ink'],P['olive'],P['orange']][index%3]
        series.graphicalProperties.line.noFill=True
    chart.set_categories(Reference(data,min_col=label_col,min_row=min_row,max_row=max_row))
    if not multi:
        chart.dataLabels=DataLabelList()
        chart.dataLabels.showVal=True
        chart.dataLabels.position='outEnd'
        chart.dataLabels.txPr=rich(8.5)
    if horizontal:
        chart.x_axis.axPos='l'
        chart.y_axis.axPos='b'
        chart.x_axis.scaling.orientation='maxMin'
        chart.y_axis.scaling.orientation='minMax'
        chart.y_axis.crosses='max'
        chart.x_axis.majorGridlines=None
        chart.y_axis.majorGridlines=ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill=P['line'],w=7000)))
    return chart

def ring(data,min_row,max_row,label_col,value_col):
    chart=chart_base(DoughnutChart())
    chart.holeSize=73
    chart.firstSliceAng=270
    chart.add_data(Reference(data,min_col=value_col,min_row=min_row-1,max_row=max_row),titles_from_data=True)
    chart.set_categories(Reference(data,min_col=label_col,min_row=min_row,max_row=max_row))
    chart.series[0].data_points=[DataPoint(idx=i,spPr=GraphicalProperties(solidFill=tone,ln=LineProperties(solidFill='FFFFFF',w=18000))) for i,tone in enumerate([P['ink'],P['olive'],P['orange']][:max_row-min_row+1])]
    # 百分比标签在正式工作簿中由办公软件原生创建，避免兼容问题。
    chart.dataLabels=None
    return chart

def line(data,min_row,max_row,label_col,value_col):
    chart=chart_base(LineChart(),legend=False)
    chart.add_data(Reference(data,min_col=value_col,min_row=min_row-1,max_row=max_row),titles_from_data=True)
    chart.set_categories(Reference(data,min_col=label_col,min_row=min_row,max_row=max_row))
    series=chart.series[0]
    series.graphicalProperties.line.solidFill=P['ink']
    series.graphicalProperties.line.width=26000
    series.marker.symbol='circle'
    series.marker.size=5
    series.marker.graphicalProperties.solidFill=P['ink']
    series.marker.graphicalProperties.line.solidFill='FFFFFF'
    return chart

def add_chart(ws,chart,region,title,subtitle,key):
    r1,c1,r2,c2=region
    panel(ws,r1,c1,r2,c2,title,subtitle)
    anchor=f'{get_column_letter(c1)}{r1+2}'
    chart.width=(c2-c1+1)*.91
    chart.height=(r2-r1-1)*.42
    ws.add_chart(chart,anchor)
    CHARTS.append({'sheet':ws.title,'key':key,'region':f'{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}','top_cell':anchor,'title':title})

def base_sheet(wb,name,title,subtitle):
    ws=wb.create_sheet(name)
    ws.sheet_view.showGridLines=False
    ws.sheet_view.showRowColHeaders=False
    ws.sheet_view.zoomScale=85
    for c in range(1,44):ws.column_dimensions[get_column_letter(c)].width=4.3
    ws.column_dimensions['A'].width=2.2
    ws.column_dimensions['AQ'].width=2.2
    for r in range(1,65):ws.row_dimensions[r].height=12
    ws.row_dimensions[2].height=18
    ws.row_dimensions[3].height=15
    ws.row_dimensions[4].height=14
    for row in ws.iter_rows(min_row=1,max_row=64,min_col=1,max_col=43):
        for cell in row:
            cell.fill=PatternFill('solid',fgColor=P['paper'])
            cell.font=Font(name='微软雅黑',size=10,color=P['ink'])
    area(ws,2,2,3,31,title,fill='ink',size=23,bold=True,color='white')
    area(ws,4,2,4,31,subtitle,fill='ink',size=9,color='white')
    area(ws,2,32,4,42,'=TEXT(TODAY(),"yyyy.mm.dd")&"  "&TEXT(TODAY(),"aaaa")',fill='ink',size=10,color='white',align='right')
    ws.print_options.horizontalCentered=True
    ws.print_options.verticalCentered=True
    ws.print_area='A1:AQ64'
    ws.page_setup.orientation='landscape'
    ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1
    ws.page_setup.fitToHeight=1
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_margins.left=.15;ws.page_margins.right=.15;ws.page_margins.top=.15;ws.page_margins.bottom=.15
    return ws

def navigation(ws,r1,c1,r2,c2,label,destination):
    cell=area(ws,r1,c1,r2,c2,f'=HYPERLINK("#\'{destination}\'!A1","{label}")',fill='white',size=11,bold=True,align='center')
    return cell

def create_component(path):
    CHARTS.clear()
    wb=Workbook()
    wb.remove(wb.active)
    home=base_sheet(wb,'经营驾驶舱','悦美 · 门店经营驾驶舱','消费趋势 / 客群结构 / 回访执行  ·  一次筛选，全页联动')
    center=base_sheet(wb,'图表分析中心','悦美 · 图表分析中心','把门店经营、服务项目与回访执行放在同一套分析视图中')
    data=wb.create_sheet('图表数据')
    for r in range(1,170):
        for c in range(1,11):data.cell(r,c,0)
    # 图表引用先占位；复制到正式工作簿后再批量写入关联公式。
    data['B2']=date.today().replace(day=1)
    for r in range(26,57):
        data.cell(r,1,r-25);data.cell(r,7,str(r-25)+'日')
    for col,title in [(2,'消费金额（元）'),(3,'已完成回访（次）')]:data.cell(25,col,title)
    for i,label in enumerate(['一店','二店','三店'],61):data.cell(i,1,label)
    data['B60']='消费金额（元）'
    for i,label in enumerate(['面部深层护理','肩颈疏通','光电嫩肤','身体护理','头皮护理','产品零售','其他'],68):data.cell(i,1,label)
    data['B67']='消费金额（元）'
    for i in range(79,88):data.cell(i,7,'员工'+str(i-78))
    data['H78']='已完成回访（次）'
    for i,label in enumerate(['周一','周二','周三','周四','周五','周六','周日'],128):data.cell(i,1,label)
    data['B127']='服务记录（次）'
    for i,label in enumerate(['1—2天','3—7天','8—14天','15天及以上'],140):data.cell(i,1,label)
    data['B139']='逾期任务（次）'
    for i in range(150,156):data.cell(i,1,str(i-149)+'月')
    data['B149']='消费金额（元）'
    for r,label in [(159,'新客'),(160,'老客'),(163,'已完成'),(164,'未完成'),(167,'已开卡'),(168,'未开卡')]:data.cell(r,1,label)
    for r in [158,162,166]:data.cell(r,2,'人数' if r!=162 else '回访次数')
    for r1,c1,r2,c2,label in [(6,2,6,4,'年份'),(6,6,6,8,'月份'),(6,10,6,16,'门店范围')]:area(home,r1,c1,r2,c2,label,fill='paper',size=9,color='muted')
    area(home,7,2,8,4,date.today().year,align='center',size=14,bold=True)
    area(home,7,6,8,8,date.today().month,align='center',size=14,bold=True,number_format='0"月"')
    area(home,7,10,8,16,'全部门店',align='center',size=12,bold=True)
    for cell,kind,f1,f2 in [('B7','whole',2025,2040),('F7','list','"1,2,3,4,5,6,7,8,9,10,11,12"',None),('J7','list','"全部门店,一店,二店,三店"',None)]:
        dv=DataValidation(type=kind,formula1=f1,formula2=f2,allow_blank=False)
        dv.showErrorMessage=True;dv.error='请使用列表中的内容或允许范围内的整数。';dv.errorTitle='筛选条件无效'
        home.add_data_validation(dv);dv.add(home[cell])
    navigation(home,6,18,8,28,'进入图表分析中心  →','图表分析中心')
    navigation(home,6,30,8,42,'查看回访月历  →','门店工作台')
    cards=[('本月消费（元）','B7','="较上月  "&IF(\'图表数据\'!B15=0,"暂无可比数据",TEXT(\'图表数据\'!B16,"+0.0%;-0.0%;0.0%"))','#,##0'),('本月到店（次）','B8','="实际服务记录  ·  按服务日统计"','0'),('在管客户（人）','B9','="新客 "&\'图表数据\'!B19&"  /  老客 "&\'图表数据\'!B20','0'),('回访完成率','B12','="本月计划 "&\'图表数据\'!B10&"  /  完成 "&\'图表数据\'!B11','0%'),('当前逾期（次）','B13','="截至今日  ·  优先处理"','0'),('生日待关怀（人）','B14','="当前待提醒  ·  提前15天"','0')]
    for i,(label,ref,note,fmt) in enumerate(cards):
        c=2+i*7
        panel(home,10,c,15,c+5,'','')
        area(home,10,c,10,c+5,label,size=9,color='muted')
        area(home,11,c,13,c+5,f"='图表数据'!{ref}",size=26,bold=True,color='red' if i==4 else 'ink',number_format=fmt)
        area(home,14,c,15,c+5,note,size=8,color='muted')
    trend=bars(data,26,56,7,[2],multi=True)
    trend.y_axis.title='消费金额（元）'
    visits=line(data,26,56,7,3)
    visits.y_axis.axId=200
    visits.y_axis.axPos='r'
    visits.y_axis.title='回访次数'
    visits.y_axis.crosses='max'
    visits.series[0].graphicalProperties.line.solidFill=P['orange']
    visits.series[0].marker.graphicalProperties.solidFill=P['orange']
    trend+=visits
    add_chart(home,trend,(18,2,36,29),'本月消费与回访走势','柱形看消费金额，折线看实际完成回访；左右坐标分别标注单位。','daily')
    add_chart(home,ring(data,159,160,1,2),(18,31,36,42),'新老客结构','按当前门店范围统计客户档案。','segments')
    add_chart(home,bars(data,61,63,1,[2],horizontal=True),(39,2,59,22),'门店消费贡献','随门店筛选联动，按服务记录消费金额汇总。','stores')
    add_chart(home,ring(data,163,164,1,2),(39,24,59,42),'本月计划执行','按计划日期统计；已停用任务不计入分母。','completion')
    area(home,61,2,62,42,'使用浅色筛选框切换年月和门店。新增服务或回访记录后，图表随公式更新。零计划时完成率显示0%，不代表存在未完成任务。',fill='paper',size=9,color='muted')
    area(center,6,2,8,29,'="当前范围："&\'经营驾驶舱\'!B7&"年"&\'经营驾驶舱\'!F7&"月  ·  "&\'经营驾驶舱\'!J7',fill='paper',size=12,bold=True)
    navigation(center,6,31,8,42,'返回驾驶舱调整筛选  →','经营驾驶舱')
    add_chart(center,line(data,150,155,1,2),(10,2,25,22),'近六个月消费走势','以所选月份为终点，观察门店消费变化。','sixmonths')
    add_chart(center,bars(data,68,74,1,[2],horizontal=True),(10,24,25,42),'服务项目消费构成','相同项目合并汇总，其余项目归入其他。','projects')
    add_chart(center,bars(data,79,87,7,[8],horizontal=True),(28,2,45,22),'员工回访排行 · 前9位','按所选月实际完成次数排序，并跟随门店范围变化。','staff')
    add_chart(center,bars(data,140,143,1,[2]),(28,24,45,42),'当前逾期时长分布','截至今日；按逾期天数分组，不随月份回溯历史。','overdue')
    add_chart(center,ring(data,167,168,1,2),(48,2,61,22),'开卡客户结构','按客户当前开卡状态统计，不与新老客分类混淆。','cards')
    add_chart(center,bars(data,128,134,1,[2]),(48,24,61,42),'周内到店节奏','按所选月实际服务日期，汇总周一至周日的到店记录。','weekdays')
    area(center,63,2,64,42,'统计口径：服务按服务日，趋势与员工回访按实际完成日，计划完成率按计划日；客户结构、逾期与生日为当前状态。',fill='paper',size=8.5,color='muted')
    data.sheet_state='hidden'
    for sheet in [home,center]:
        merges=list(sheet.merged_cells.ranges)
        for i,a in enumerate(merges):
            for b in merges[i+1:]:
                assert a.max_row<b.min_row or b.max_row<a.min_row or a.max_col<b.min_col or b.max_col<a.min_col,(sheet.title,str(a),str(b),'合并区域不能重叠')
    # 组件仅传递页面，图表由办公软件在正式工作簿中原生创建。
    # 这样可以避免跨文件图表对象的兼容问题和临时文件链接。
    home._charts=[]
    center._charts=[]
    wb.active=0
    wb.save(path)
    return list(CHARTS)
